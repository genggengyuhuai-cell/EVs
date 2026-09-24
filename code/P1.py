from pathlib import Path
import re
from collections import Counter, defaultdict

import pandas as pd
from openpyxl import load_workbook


# ============================================================
# 1. Paths
# ============================================================

INPUT_FILE = Path("../rawdata/processed.xlsx")
OUTPUT_FILE = Path("../rawdata/sample_mapping_audit.xlsx")


# ============================================================
# 2. Helper functions
# ============================================================

def clean_text(x):
    if pd.isna(x):
        return ""

    x = str(x).strip()

    if x.endswith(".0"):
        x = x[:-2]

    return x


def normalize_id(x):
    """
    Normalize separators only.
    IMPORTANT:
    Does NOT remove biologically meaningful group information.
    """
    x = clean_text(x)

    x = x.replace(" ", "")

    # collapse repeated underscores
    x = re.sub(
        r"_+",
        "_",
        x
    )

    # remove leading/trailing underscore
    x = x.strip("_")

    return x


def clean_treat1(x):
    x = clean_text(x).lower()

    mapping = {
        "low": "low",
        "high": "high",
        "contrl": "control",
        "control": "control",
        "unknown": "unknown",
        "": "missing"
    }

    return mapping.get(
        x,
        x
    )


# ============================================================
# 3. Read ORIGINAL Sheet1 headers using openpyxl
#
# IMPORTANT:
# pandas automatically converts duplicate columns into
# xxx, xxx.1, xxx.2.
#
# openpyxl lets us see the REAL original Excel headers.
# ============================================================

wb = load_workbook(
    INPUT_FILE,
    read_only=True,
    data_only=True
)

ws1 = wb.worksheets[0]

header_row = next(
    ws1.iter_rows(
        min_row=1,
        max_row=1,
        values_only=True
    )
)

all_headers = [
    clean_text(x)
    for x in header_row
]

annotation_headers = all_headers[:7]

sample_headers_raw = all_headers[7:]

print(
    "============================================================"
)
print(
    "SHEET1 ORIGINAL HEADER CHECK"
)
print(
    "============================================================"
)

print(
    "Annotation columns :",
    len(annotation_headers)
)

print(
    "Sample columns     :",
    len(sample_headers_raw)
)

print(
    "\nAnnotation columns:"
)

for x in annotation_headers:
    print(
        "  ",
        x
    )


# ============================================================
# 4. Read metadata
# ============================================================

meta = pd.read_excel(
    INPUT_FILE,
    sheet_name=1
).copy()

required_columns = [
    "进样时间",
    "sample",
    "condition",
    "group",
    "TREAT1"
]

missing_columns = [
    x
    for x in required_columns
    if x not in meta.columns
]

if missing_columns:
    raise ValueError(
        f"Missing metadata columns: {missing_columns}"
    )


meta["metadata_row"] = (
    meta.index +
    2
)

meta["date_clean"] = (
    meta["进样时间"]
    .apply(clean_text)
)

meta["sample_clean"] = (
    meta["sample"]
    .apply(clean_text)
)

meta["group_clean"] = (
    meta["group"]
    .apply(clean_text)
)

meta["condition_clean"] = (
    meta["condition"]
    .apply(clean_text)
)

meta["TREAT1_clean"] = (
    meta["TREAT1"]
    .apply(clean_treat1)
)


# ============================================================
# 5. Construct STRICT unique metadata ID
# ============================================================

meta["UniqueSampleID"] = [
    normalize_id(
        f"{date}_{group}_{sample}"
    )
    for date, group, sample in zip(
        meta["date_clean"],
        meta["group_clean"],
        meta["sample_clean"]
    )
]


# ============================================================
# 6. Construct possible aliases
#
# Example:
#
# Metadata:
# date   = 20251001
# group  = XZ_YA
# sample = 1
#
# aliases:
# 20251001_1
# 20251001_XZ_YA_1
# 20251001_XZ_1
#
#
# Metadata:
# group  = FJ_FQ
# sample = F83
#
# aliases:
# 20251001_F83
# 20251001_FJ_FQ_F83
# 20251001_FJ_F83
#
# IMPORTANT:
# aliases are used only to generate candidates.
# Ambiguous candidates are NEVER automatically assigned.
# ============================================================

metadata_aliases = {}

alias_to_metadata_rows = defaultdict(
    list
)

for idx, row in meta.iterrows():

    date = row["date_clean"]
    sample = row["sample_clean"]
    group = row["group_clean"]

    aliases = set()

    # date + sample
    aliases.add(
        normalize_id(
            f"{date}_{sample}"
        )
    )

    # date + full group + sample
    aliases.add(
        normalize_id(
            f"{date}_{group}_{sample}"
        )
    )

    # first group component:
    # FJ_FQ -> FJ
    # XZ_YA -> XZ
    # GZ_TH -> GZ
    if group:

        group_prefix = group.split("_")[0]

        aliases.add(
            normalize_id(
                f"{date}_{group_prefix}_{sample}"
            )
        )

    metadata_aliases[idx] = aliases

    for alias in aliases:

        alias_to_metadata_rows[
            alias
        ].append(
            idx
        )


# ============================================================
# 7. Normalize Sheet1 ORIGINAL sample headers
# ============================================================

sheet1_table = pd.DataFrame({
    "Sheet1_position": range(
        1,
        len(sample_headers_raw) + 1
    ),
    "Excel_column": range(
        8,
        len(sample_headers_raw) + 8
    ),
    "Sheet1_raw_header": sample_headers_raw
})

sheet1_table[
    "Sheet1_normalized"
] = (
    sheet1_table[
        "Sheet1_raw_header"
    ]
    .apply(
        normalize_id
    )
)


# ============================================================
# 8. Check REAL duplicate headers
# ============================================================

header_counts = Counter(
    sheet1_table[
        "Sheet1_normalized"
    ]
)

sheet1_table[
    "Sheet1_duplicate_count"
] = (
    sheet1_table[
        "Sheet1_normalized"
    ]
    .map(
        header_counts
    )
)


# ============================================================
# 9. Generate matching candidates
# ============================================================

candidate_records = []

for _, srow in sheet1_table.iterrows():

    sid = srow[
        "Sheet1_normalized"
    ]

    candidate_indices = (
        alias_to_metadata_rows
        .get(
            sid,
            []
        )
    )

    candidate_records.append({
        "Sheet1_position":
            srow["Sheet1_position"],

        "Excel_column":
            srow["Excel_column"],

        "Sheet1_raw_header":
            srow["Sheet1_raw_header"],

        "Sheet1_normalized":
            sid,

        "Sheet1_duplicate_count":
            srow["Sheet1_duplicate_count"],

        "candidate_count":
            len(candidate_indices),

        "candidate_metadata_indices":
            candidate_indices
    })

candidate_df = pd.DataFrame(
    candidate_records
)


# ============================================================
# 10. Initial classification
#
# VERY IMPORTANT:
#
# UNIQUE_NAME_MATCH:
# only one metadata row matches the name
#
# AMBIGUOUS:
# >1 metadata rows possible
#
# UNMATCHED:
# no metadata row possible
#
# But even UNIQUE_NAME_MATCH is checked globally below.
# ============================================================

candidate_df[
    "Initial_status"
] = candidate_df[
    "candidate_count"
].map(
    lambda x:
        "UNMATCHED"
        if x == 0
        else (
            "UNIQUE_NAME_MATCH"
            if x == 1
            else "AMBIGUOUS"
        )
)


# ============================================================
# 11. Detect cases where multiple Sheet1 columns point to
#     the SAME metadata row
#
# A one-to-one mapping must satisfy BOTH:
#
# Sheet1 column -> one metadata row
# metadata row  -> one Sheet1 column
# ============================================================

metadata_target_counts = Counter()

for candidates in candidate_df[
    "candidate_metadata_indices"
]:

    if len(candidates) == 1:

        metadata_target_counts[
            candidates[0]
        ] += 1


final_records = []

for _, row in candidate_df.iterrows():

    candidates = row[
        "candidate_metadata_indices"
    ]

    status = row[
        "Initial_status"
    ]

    matched_idx = None

    if len(candidates) == 1:

        possible_idx = candidates[0]

        if metadata_target_counts[
            possible_idx
        ] == 1:

            status = "CONFIRMED_UNIQUE"

            matched_idx = possible_idx

        else:

            status = (
                "CONFLICT_MULTIPLE_COLUMNS_TO_ONE_METADATA"
            )

    record = row.to_dict()

    record[
        "Final_status"
    ] = status

    if matched_idx is not None:

        mrow = meta.loc[
            matched_idx
        ]

        record[
            "metadata_row"
        ] = mrow[
            "metadata_row"
        ]

        record[
            "UniqueSampleID"
        ] = mrow[
            "UniqueSampleID"
        ]

        record[
            "进样时间"
        ] = mrow[
            "进样时间"
        ]

        record[
            "sample"
        ] = mrow[
            "sample"
        ]

        record[
            "condition"
        ] = mrow[
            "condition"
        ]

        record[
            "group"
        ] = mrow[
            "group"
        ]

        record[
            "TREAT1"
        ] = mrow[
            "TREAT1"
        ]

        record[
            "TREAT1_clean"
        ] = mrow[
            "TREAT1_clean"
        ]

        if "TREAT2" in meta.columns:

            record[
                "TREAT2"
            ] = mrow[
                "TREAT2"
            ]

    else:

        record[
            "metadata_row"
        ] = None

        record[
            "UniqueSampleID"
        ] = None

        record[
            "进样时间"
        ] = None

        record[
            "sample"
        ] = None

        record[
            "condition"
        ] = None

        record[
            "group"
        ] = None

        record[
            "TREAT1"
        ] = None

        record[
            "TREAT1_clean"
        ] = None

        # print candidate details for manual audit
        descriptions = []

        for idx in candidates:

            mrow = meta.loc[
                idx
            ]

            descriptions.append(
                " | ".join([
                    f"ExcelRow={mrow['metadata_row']}",
                    f"ID={mrow['UniqueSampleID']}",
                    f"condition={mrow['condition_clean']}",
                    f"group={mrow['group_clean']}",
                    f"TREAT1={mrow['TREAT1_clean']}"
                ])
            )

        record[
            "Candidate_details"
        ] = " || ".join(
            descriptions
        )

    final_records.append(
        record
    )


mapping = pd.DataFrame(
    final_records
)


# ============================================================
# 12. Find metadata rows that have NOT been uniquely mapped
# ============================================================

confirmed_metadata_rows = set(
    mapping.loc[
        mapping[
            "Final_status"
        ] == "CONFIRMED_UNIQUE",
        "metadata_row"
    ]
    .dropna()
    .astype(int)
)


unmapped_metadata = meta.loc[
    ~meta[
        "metadata_row"
    ].isin(
        confirmed_metadata_rows
    )
].copy()


# ============================================================
# 13. Metadata uniqueness QC
# ============================================================

unique_id_counts = (
    meta[
        "UniqueSampleID"
    ]
    .value_counts()
)

duplicate_unique_ids = (
    unique_id_counts[
        unique_id_counts > 1
    ]
    .reset_index()
)

duplicate_unique_ids.columns = [
    "UniqueSampleID",
    "Count"
]


# ============================================================
# 14. Summary
# ============================================================

status_counts = (
    mapping[
        "Final_status"
    ]
    .value_counts()
)

summary_rows = [
    [
        "Sheet1 sample columns",
        len(sample_headers_raw)
    ],
    [
        "Sheet2 metadata rows",
        len(meta)
    ],
    [
        "CONFIRMED_UNIQUE",
        int(
            status_counts.get(
                "CONFIRMED_UNIQUE",
                0
            )
        )
    ],
    [
        "AMBIGUOUS",
        int(
            status_counts.get(
                "AMBIGUOUS",
                0
            )
        )
    ],
    [
        "UNMATCHED",
        int(
            status_counts.get(
                "UNMATCHED",
                0
            )
        )
    ],
    [
        "CONFLICT_MULTIPLE_COLUMNS_TO_ONE_METADATA",
        int(
            status_counts.get(
                "CONFLICT_MULTIPLE_COLUMNS_TO_ONE_METADATA",
                0
            )
        )
    ],
    [
        "Metadata rows not uniquely mapped",
        len(
            unmapped_metadata
        )
    ],
    [
        "Duplicate UniqueSampleID in metadata",
        len(
            duplicate_unique_ids
        )
    ]
]

summary = pd.DataFrame(
    summary_rows,
    columns=[
        "Metric",
        "Value"
    ]
)


# ============================================================
# 15. Specific design summaries
# ============================================================

condition_counts = (
    meta[
        "condition_clean"
    ]
    .value_counts(
        dropna=False
    )
    .rename_axis(
        "condition"
    )
    .reset_index(
        name="N"
    )
)

group_counts = (
    meta[
        "group_clean"
    ]
    .value_counts(
        dropna=False
    )
    .rename_axis(
        "group"
    )
    .reset_index(
        name="N"
    )
)

treat_counts = (
    meta[
        "TREAT1_clean"
    ]
    .value_counts(
        dropna=False
    )
    .rename_axis(
        "TREAT1"
    )
    .reset_index(
        name="N"
    )
)


# ============================================================
# 16. Write audit workbook
# ============================================================

with pd.ExcelWriter(
    OUTPUT_FILE,
    engine="openpyxl"
) as writer:

    summary.to_excel(
        writer,
        sheet_name="00_Summary",
        index=False
    )

    mapping.to_excel(
        writer,
        sheet_name="01_All_Mapping",
        index=False
    )

    mapping.loc[
        mapping[
            "Final_status"
        ] == "CONFIRMED_UNIQUE"
    ].to_excel(
        writer,
        sheet_name="02_Confirmed",
        index=False
    )

    mapping.loc[
        mapping[
            "Final_status"
        ] != "CONFIRMED_UNIQUE"
    ].to_excel(
        writer,
        sheet_name="03_Need_Review",
        index=False
    )

    unmapped_metadata.to_excel(
        writer,
        sheet_name="04_Unmapped_Metadata",
        index=False
    )

    sheet1_table.to_excel(
        writer,
        sheet_name="05_Sheet1_Headers",
        index=False
    )

    duplicate_unique_ids.to_excel(
        writer,
        sheet_name="06_Duplicate_Metadata_ID",
        index=False
    )

    condition_counts.to_excel(
        writer,
        sheet_name="07_Condition_Counts",
        index=False
    )

    group_counts.to_excel(
        writer,
        sheet_name="08_Group_Counts",
        index=False
    )

    treat_counts.to_excel(
        writer,
        sheet_name="09_TREAT1_Counts",
        index=False
    )


# ============================================================
# 17. Console report
# ============================================================

print(
    "\n============================================================"
)
print(
    "SAMPLE MAPPING AUDIT COMPLETED"
)
print(
    "============================================================"
)

print(
    summary.to_string(
        index=False
    )
)

print(
    "\nOutput:"
)

print(
    OUTPUT_FILE
)

print(
    "\nIMPORTANT:"
)

if (
    len(mapping) == len(meta)
    and
    (
        mapping[
            "Final_status"
        ] == "CONFIRMED_UNIQUE"
    ).all()
    and
    len(unmapped_metadata) == 0
    and
    len(duplicate_unique_ids) == 0
):

    print(
        "PASS: all 519 samples are uniquely and safely mapped."
    )

else:

    print(
        "NOT YET PASS."
    )

    print(
        "Do NOT use this mapping for proteomics analysis yet."
    )

    print(
        "Open sheet '03_Need_Review' in sample_mapping_audit.xlsx."
    )

    print(
        "Ambiguous/unmatched samples must be resolved first."
    )