from pathlib import Path
import re

import pandas as pd
from openpyxl import load_workbook


# ============================================================
# 1. Paths
# ============================================================

INPUT_FILE = Path("../rawdata/processed.xlsx")
AUDIT_FILE = Path("../rawdata/sample_mapping_audit.xlsx")

OUTPUT_FILE = Path("../rawdata/sample_mapping_FINAL.xlsx")


# ============================================================
# 2. Helpers
# ============================================================

def clean_text(x):

    if pd.isna(x):
        return ""

    x = str(x).strip()

    if x.endswith(".0"):
        x = x[:-2]

    return x


def normalize_id(x):

    x = clean_text(x)

    x = re.sub(
        r"_+",
        "_",
        x
    )

    return x.strip("_")


def clean_treat1(x):

    x = clean_text(
        x
    ).lower()

    mapping = {
        "contrl": "control",
        "control": "control",
        "low": "low",
        "high": "high",
        "unknown": "unknown",
        "": "missing"
    }

    return mapping.get(
        x,
        x
    )


# ============================================================
# 3. Read original Sheet1 headers
# ============================================================

wb = load_workbook(
    INPUT_FILE,
    read_only=True,
    data_only=True
)

ws = wb.worksheets[0]

header = next(
    ws.iter_rows(
        min_row=1,
        max_row=1,
        values_only=True
    )
)

sample_headers = [
    clean_text(x)
    for x in header[7:]
]


# ============================================================
# 4. Read metadata
# ============================================================

meta = pd.read_excel(
    INPUT_FILE,
    sheet_name=1
).copy()

meta["metadata_index"] = meta.index

meta["metadata_excel_row"] = (
    meta.index +
    2
)

meta["date_clean"] = (
    meta["进样时间"]
    .apply(clean_text)
)

meta["sample_clean"] = (
    meta["sample"]
    .apply(clean_text)
)

meta["group_clean"] = (
    meta["group"]
    .apply(clean_text)
)

meta["TREAT1_clean"] = (
    meta["TREAT1"]
    .apply(clean_treat1)
)

meta["UniqueSampleID"] = [
    normalize_id(
        f"{date}_{group}_{sample}"
    )
    for date, group, sample in zip(
        meta["date_clean"],
        meta["group_clean"],
        meta["sample_clean"]
    )
]


# ============================================================
# 5. Read P1 confirmed mapping
# ============================================================

old = pd.read_excel(
    AUDIT_FILE,
    sheet_name="01_All_Mapping"
)

if len(old) != 519:
    raise ValueError(
        f"P1 mapping rows = {len(old)}, expected 519."
    )


# ============================================================
# 6. Build final mapping container
# ============================================================

records = []

used_metadata = set()


# ============================================================
# 7. First retain all previously CONFIRMED_UNIQUE samples
# ============================================================

for _, row in old.iterrows():

    pos = int(
        row["Sheet1_position"]
    )

    header_raw = sample_headers[
        pos - 1
    ]

    record = {
        "Sheet1_position": pos,
        "Sheet1_raw_header": header_raw,
        "Resolution_method": None,
        "metadata_index": None
    }

    if (
        row["Final_status"]
        ==
        "CONFIRMED_UNIQUE"
    ):

        metadata_row = int(
            row["metadata_row"]
        )

        metadata_index = (
            metadata_row -
            2
        )

        record[
            "metadata_index"
        ] = metadata_index

        record[
            "Resolution_method"
        ] = "P1_UNIQUE_NAME"

        used_metadata.add(
            metadata_index
        )

    records.append(
        record
    )


final = pd.DataFrame(
    records
)


# ============================================================
# 8. Resolve known ambiguous contiguous blocks
#
# These blocks were established from:
#
# 1) neighboring confirmed group boundaries
# 2) exact group-specific sample number sets
# 3) continuous Sheet1 ordering
#
# Sheet1_position is 1-based among the 519 sample columns.
# ============================================================

block_rules = [
    {
        "start": 166,
        "end": 177,
        "date": "20251001",
        "group": "XZ_YB"
    },
    {
        "start": 178,
        "end": 205,
        "date": "20251001",
        "group": "XZ_YC"
    },
    {
        "start": 212,
        "end": 238,
        "date": "20251001",
        "group": "XZ_YD"
    }
]

# ============================================================
# 9. Resolve each block by SAMPLE NUMBER, not merely row order
# ============================================================

for rule in block_rules:

    block_positions = list(
        range(
            rule["start"],
            rule["end"] + 1
        )
    )

    candidates = meta.loc[
        (
            meta["date_clean"]
            ==
            rule["date"]
        )
        &
        (
            meta["group_clean"]
            ==
            rule["group"]
        )
    ].copy()

    candidate_by_sample = {
        str(row["sample_clean"]):
            int(row["metadata_index"])
        for _, row in candidates.iterrows()
    }

    for pos in block_positions:

        raw_header = sample_headers[
            pos - 1
        ]

        normalized = normalize_id(
            raw_header
        )

        # extract the final numeric sample component
        match = re.search(
            r"_(\d+)$",
            normalized
        )

        if match is None:
            raise ValueError(
                f"Cannot extract sample number "
                f"from position {pos}: {raw_header}"
            )

        sample_number = match.group(
            1
        )

        if (
            sample_number
            not in
            candidate_by_sample
        ):
            raise ValueError(
                f"Position {pos}: sample {sample_number} "
                f"does not exist in "
                f"{rule['date']} {rule['group']} metadata."
            )

        metadata_index = (
            candidate_by_sample[
                sample_number
            ]
        )

        if metadata_index in used_metadata:

            raise ValueError(
                f"Metadata index {metadata_index} "
                f"would be assigned twice."
            )

        final.loc[
            final["Sheet1_position"] == pos,
            "metadata_index"
        ] = metadata_index

        final.loc[
            final["Sheet1_position"] == pos,
            "Resolution_method"
        ] = (
            "BLOCK_BOUNDARY_AND_SAMPLE_ID"
        )

        used_metadata.add(
            metadata_index
        )


# ============================================================
# 10. Check whether anything remains unresolved
# ============================================================

unresolved = final.loc[
    final["metadata_index"].isna()
].copy()

if len(unresolved) > 0:

    print(
        "\nUNRESOLVED SHEET1 COLUMNS:"
    )

    print(
        unresolved.to_string(
            index=False
        )
    )

    raise ValueError(
        f"{len(unresolved)} Sheet1 columns remain unresolved."
    )


# ============================================================
# 11. Force integer index after complete resolution
# ============================================================

final["metadata_index"] = (
    final["metadata_index"]
    .astype(int)
)


# ============================================================
# 12. Attach metadata
# ============================================================

meta_lookup = (
    meta
    .set_index(
        "metadata_index"
    )
)

metadata_columns = [
    "metadata_excel_row",
    "UniqueSampleID",
    "进样时间",
    "sample",
    "condition",
    "group",
    "TREAT1",
    "TREAT1_clean",
    "TREAT2",
    "Tube_Mixing",
    "WoleBlood_oldTime",
    "Plasma_HoldTime_h",
    "Note"
]

for col in metadata_columns:

    if col in meta_lookup.columns:

        final[col] = [
            meta_lookup.loc[
                idx,
                col
            ]
            for idx in final[
                "metadata_index"
            ]
        ]


# ============================================================
# 13. HARD validation
# ============================================================

n_sheet1 = len(
    final
)

n_metadata = len(
    meta
)

n_unique_metadata_assignments = (
    final[
        "metadata_index"
    ]
    .nunique()
)

n_unique_ids = (
    final[
        "UniqueSampleID"
    ]
    .nunique()
)

duplicate_mapping = (
    final[
        "metadata_index"
    ]
    .duplicated(
        keep=False
    )
)

missing_metadata_indices = sorted(
    set(
        meta[
            "metadata_index"
        ]
    )
    -
    set(
        final[
            "metadata_index"
        ]
    )
)


# ============================================================
# 14. Validate final group/date sample sets
#
# IMPORTANT:
# Validate the FULL final mapping for each group,
# not only the ambiguous sub-block.
# ============================================================

groups_to_validate = [
    {
        "date": "20251001",
        "group": "XZ_YB"
    },
    {
        "date": "20251001",
        "group": "XZ_YC"
    },
    {
        "date": "20251001",
        "group": "XZ_YD"
    }
]

block_validation = []

for rule in groups_to_validate:

    observed = final.loc[
        (
            final["进样时间"]
            .astype(str)
            .str.replace(r"\.0$", "", regex=True)
            ==
            rule["date"]
        )
        &
        (
            final["group"]
            ==
            rule["group"]
        ),
        "sample"
    ].astype(
        str
    ).tolist()

    expected = meta.loc[
        (
            meta["date_clean"]
            ==
            rule["date"]
        )
        &
        (
            meta["group_clean"]
            ==
            rule["group"]
        ),
        "sample_clean"
    ].astype(
        str
    ).tolist()

    observed_set = set(
        observed
    )

    expected_set = set(
        expected
    )

    block_validation.append({
        "Date":
            rule["date"],

        "Group":
            rule["group"],

        "Observed_N":
            len(observed),

        "Expected_N":
            len(expected),

        "Sample_sets_identical":
            observed_set
            ==
            expected_set,

        "Missing_from_final_mapping":
            ",".join(
                sorted(
                    expected_set
                    -
                    observed_set
                )
            ),

        "Unexpected_in_final_mapping":
            ",".join(
                sorted(
                    observed_set
                    -
                    expected_set
                )
            )
    })

block_validation = pd.DataFrame(
    block_validation
)


# ============================================================
# 15. Summary
# ============================================================

summary = pd.DataFrame({
    "Metric": [
        "Sheet1 sample columns",
        "Sheet2 metadata rows",
        "Mapped rows",
        "Unique metadata assignments",
        "Unique final SampleIDs",
        "Duplicated metadata assignments",
        "Missing metadata rows",
        "P1 unique-name mappings",
        "Block-resolved mappings"
    ],
    "Value": [
        n_sheet1,
        n_metadata,
        len(final),
        n_unique_metadata_assignments,
        n_unique_ids,
        int(
            duplicate_mapping.sum()
        ),
        len(
            missing_metadata_indices
        ),
        int(
            (
                final[
                    "Resolution_method"
                ]
                ==
                "P1_UNIQUE_NAME"
            ).sum()
        ),
        int(
            (
                final[
                    "Resolution_method"
                ]
                ==
                "BLOCK_BOUNDARY_AND_SAMPLE_ID"
            ).sum()
        )
    ]
})


# ============================================================
# 16. Final PASS criteria
# ============================================================

blocks_pass = (
    block_validation[
        "Sample_sets_identical"
    ].all()
)

mapping_pass = (
    n_sheet1
    ==
    519
    and
    n_metadata
    ==
    519
    and
    n_unique_metadata_assignments
    ==
    519
    and
    n_unique_ids
    ==
    519
    and
    duplicate_mapping.sum()
    ==
    0
    and
    len(
        missing_metadata_indices
    )
    ==
    0
    and
    blocks_pass
)


# ============================================================
# 17. Save FINAL audit workbook
# ============================================================

with pd.ExcelWriter(
    OUTPUT_FILE,
    engine="openpyxl"
) as writer:

    summary.to_excel(
        writer,
        sheet_name="00_FINAL_SUMMARY",
        index=False
    )

    final.to_excel(
        writer,
        sheet_name="01_FINAL_MAPPING",
        index=False
    )

    block_validation.to_excel(
        writer,
        sheet_name="02_BLOCK_VALIDATION",
        index=False
    )

    meta.to_excel(
        writer,
        sheet_name="03_METADATA_REFERENCE",
        index=False
    )


# ============================================================
# 18. Console output
# ============================================================

print(
    "\n============================================================"
)

print(
    "FINAL SAMPLE MAPPING VALIDATION"
)

print(
    "============================================================"
)

print(
    summary.to_string(
        index=False
    )
)

print(
    "\nBLOCK VALIDATION:"
)

print(
    block_validation.to_string(
        index=False
    )
)

print(
    "\n============================================================"
)

if mapping_pass:

    print(
        "PASS: ALL 519 SAMPLES ARE UNIQUELY MAPPED."
    )

    print(
        "This mapping can be locked for downstream analysis."
    )

else:

    print(
        "FAIL: mapping is not yet safe."
    )

    print(
        "Do NOT proceed to proteomics analysis."
    )

print(
    "\nOutput:"
)

print(
    OUTPUT_FILE
)

print(
    "============================================================"
)