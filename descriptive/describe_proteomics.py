from pathlib import Path
import hashlib
import json
import numpy as np
import pandas as pd
from openpyxl import load_workbook
import matplotlib
matplotlib.use('Agg')
from matplotlib.colors import ListedColormap
from nature_plotting import new_figure, save_series

matplotlib.rcParams["font.sans-serif"] = [
    "Microsoft YaHei", "SimHei", "Arial", "DejaVu Sans"
]

ROOT = Path(__file__).resolve().parents[1]
OUT = Path(__file__).resolve().parent
SOURCE = ROOT / 'rawdata' / 'processed.xlsx'
MAP = ROOT / 'rawdata' / 'sample_mapping_FINAL.xlsx'

def clean(v):
    if pd.isna(v): return ''
    s = str(v).strip()
    return s[:-2] if s.endswith('.0') else s

wb = load_workbook(SOURCE, read_only=True, data_only=True)
rows = wb.worksheets[0].iter_rows(values_only=True)
headers = next(rows)
data = pd.DataFrame(list(rows), columns=range(len(headers)))
wb.close()
annotation = data.iloc[:,:7].copy()
annotation.columns = headers[:7]
raw = data.iloc[:,7:]
values = raw.apply(pd.to_numeric,errors='coerce').to_numpy(dtype=float)
meta = pd.read_excel(SOURCE,sheet_name=1)
mapping = pd.read_excel(MAP,sheet_name='01_FINAL_MAPPING').sort_values('Sheet1_position').reset_index(drop=True)
nprot, nsamp = values.shape
assert len(mapping)==len(meta)==nsamp
assert mapping.Sheet1_position.tolist()==list(range(1,nsamp+1))
assert mapping.metadata_index.nunique()==nsamp and mapping.UniqueSampleID.nunique()==nsamp
assert set(mapping.metadata_index)==set(range(len(meta)))
assert [clean(x) for x in headers[7:]]==mapping.Sheet1_raw_header.map(clean).tolist()
for col in meta.columns:
    assert col in mapping.columns, col
    expected=meta.iloc[mapping.metadata_index.astype(int)][col].map(clean).tolist()
    assert expected==mapping[col].map(clean).tolist(), col
assert annotation.iloc[:,0].notna().all() and not annotation.iloc[:,0].duplicated().any()
detected=np.isfinite(values)&(values>0)
missing=~detected
non_numeric=int((raw.notna().to_numpy() & np.isnan(values)).sum())
audit={'mapping_check':'PASS: headers, positions, metadata fields and bijection checked','protein_group_rows':nprot,'samples':nsamp,'total_cells':int(values.size),'original_null_cells':int(raw.isna().sum().sum()),'non_numeric_cells':non_numeric,'zero_cells':int((values==0).sum()),'negative_cells':int((values<0).sum()),'infinite_cells':int(np.isinf(values).sum()),'missing_cells':int(missing.sum()),'missing_pct':float(missing.mean()*100),'detected_at_least_once':int(detected.any(axis=1).sum()),'detected_in_all_samples':int(detected.all(axis=1).sum()),'never_detected':int((~detected.any(axis=1)).sum())}
for p in [SOURCE,MAP]: audit[p.name+'_sha256']=hashlib.sha256(p.read_bytes()).hexdigest()
(OUT/'audit.json').write_text(json.dumps(audit,ensure_ascii=False,indent=2),encoding='utf-8')
sample=mapping.copy()
sample['detected_protein_groups']=detected.sum(axis=0)
sample['missing_pct']=missing.mean(axis=0)*100
sample.to_csv(OUT/'sample_statistics.csv',index=False,encoding='utf-8-sig')
protein=annotation.copy()
protein['detected_samples']=detected.sum(axis=1)
protein['missing_pct']=missing.mean(axis=1)*100
positive=np.where(detected,values,np.nan)
protein['median_positive_abundance']=pd.DataFrame(positive).median(axis=1).to_numpy()
protein.to_csv(OUT/'protein_statistics.csv',index=False,encoding='utf-8-sig')

def summarize(field):
    records=[]
    prevalence=pd.DataFrame({'PG.ProteinGroups':annotation.iloc[:,0]})
    for name, sub in sample.groupby(field,dropna=False,sort=True):
        ix=sub.index.to_numpy(); d=detected[:,ix]; n=len(ix)
        rates=d.mean(axis=1)
        prevalence[str(name)]=rates
        records.append({'category':str(name),'samples':n,'detected_any':int((rates>0).sum()),'detected_ge50pct':int((rates>=0.5).sum()),'detected_ge80pct':int((rates>=0.8).sum()),'detected_all':int((rates==1).sum()),'missing_pct':float((~d).mean()*100),'sample_detected_median':float(np.median(d.sum(axis=0))),'sample_detected_q25':float(np.quantile(d.sum(axis=0),.25)),'sample_detected_q75':float(np.quantile(d.sum(axis=0),.75)),'sample_detected_min':int(d.sum(axis=0).min()),'sample_detected_max':int(d.sum(axis=0).max())})
    result=pd.DataFrame(records)
    result.to_csv(OUT/f'{field}_summary.csv',index=False,encoding='utf-8-sig')
    prevalence.to_csv(OUT/f'{field}_protein_detection_rate.csv',index=False,encoding='utf-8-sig')
    return result,prevalence

env,ep=summarize('condition'); region,rp=summarize('group')
summarize('TREAT1_clean'); summarize('进样时间')
cross=pd.crosstab(sample['group'],sample['condition'])
cross.to_csv(OUT/'region_environment_counts.csv',encoding='utf-8-sig')
e1=ep.iloc[:,1].to_numpy()>0; e2=ep.iloc[:,2].to_numpy()>0
overlap={'shared':int((e1&e2).sum()),str(ep.columns[1])+'_only':int((e1&~e2).sum()),str(ep.columns[2])+'_only':int((e2&~e1).sum())}
COLORS={'high_stress':'#527D9E','high_temperature':'#C48C59',
        '高海拔':'#527D9E','湿热':'#C48C59'}
LABELS={'high_stress':'高海拔','high_temperature':'湿热',
        '高海拔':'高海拔','湿热':'湿热'}
def label(ax,title):
    ax.set_title(title,loc='left',pad=9)

coverage_figures, coverage_axes = zip(*(new_figure() for _ in range(4)))
axs = np.array(coverage_axes, dtype=object).reshape(2, 2)
ax=axs[0,0]; y=np.arange(len(region)); colors=[COLORS[sample.loc[sample['group']==g,'condition'].iloc[0]] for g in region.category]
ax.barh(y,region.samples,color=colors)
ax.set_yticks(y,region.category); ax.invert_yaxis(); ax.set_xlim(0,region.samples.max()*1.2); ax.set_xlabel('Samples (n)')
for i,n in enumerate(region.samples): ax.text(n+2,i,str(n),va='center',fontsize=6)
label(ax,'Cohort composition')
ax=axs[0,1]; x=np.arange(len(env)); width=.23
for j,(col,name,c) in enumerate([('detected_any','At least 1 sample','#B8C4CE'),('detected_ge80pct','At least 80%','#527D9E'),('detected_all','All samples','#283D4D')]):
    bars=ax.bar(x+(j-1)*width,env[col],width,color=c,label=name)
    ax.bar_label(bars,padding=2,fontsize=6,rotation=90)
ax.set_xticks(x,[LABELS[z]+f'\n(n={n})' for z,n in zip(env.category,env.samples)]); ax.set_ylim(0,nprot*1.35); ax.set_ylabel('Detected protein groups'); ax.legend(loc='upper left',frameon=False)
label(ax,'Coverage by environment')
ax=axs[1,0]
for col,name,c in [('detected_any','At least 1 sample','#B8C4CE'),('detected_ge80pct','At least 80%','#527D9E'),('detected_all','All samples','#283D4D')]:
    ax.plot(region[col],y,'o',ms=4,color=c,label=name)
ax.set_yticks(y,region.category); ax.invert_yaxis(); ax.set_xlim(0,nprot*1.05); ax.set_xlabel('Detected protein groups')
label(ax,'Coverage by region')
ax=axs[1,1]; groups=[sample.loc[sample['group']==g,'detected_protein_groups'].to_numpy() for g in region.category]
bp=ax.boxplot(groups,vert=False,patch_artist=True,widths=.55,showfliers=True,flierprops={'markersize':2},medianprops={'color':'black','linewidth':.9})
for box,c in zip(bp['boxes'],colors): box.set_facecolor(c)
ax.set_yticks(np.arange(1,len(region)+1),region.category); ax.invert_yaxis(); ax.set_xlabel('Detected protein groups per sample'); ax.set_xlim(0,nprot*1.05)
label(ax,'Sample-level depth')
save_series(coverage_figures, OUT, ['Figure_01_descriptive_protein_coverage', 'Figure_01_descriptive_environment_coverage',
                                   'Figure_01_descriptive_region_coverage', 'Figure_01_descriptive_sample_depth'],
            [region[['category','samples']], env, region,
             sample[['UniqueSampleID','group','condition','detected_protein_groups']]])

missing_figures, missing_axes = zip(*(new_figure(183, 140) for _ in range(3)))
ax=missing_axes[0]; ordered=sample.sort_values(['condition','group','Sheet1_position']).index.to_numpy(); orderp=np.argsort(missing.mean(axis=1),kind='stable')
ax.imshow(missing[orderp][:,ordered],aspect='auto',interpolation='nearest',cmap=ListedColormap(['#E4EBEF','#344F65']),rasterized=True,vmin=0,vmax=1)
ticks=[]; names=[]; last=0
for (condition,group),sub in sample.loc[ordered].groupby(['condition','group'],sort=False):
    ticks.append(last+(len(sub)-1)/2); names.append(group)
    if last: ax.axvline(last-.5,color='white',linewidth=.6)
    last+=len(sub)
ax.set_xticks(ticks,names,rotation=45,ha='right',rotation_mode='anchor'); ax.set_ylabel('Protein groups\n(sorted by missingness)'); ax.set_xlabel('All samples, ordered by environment and region')
label(ax,'Missingness across the full matrix (dark = missing)')
ax=missing_axes[1]; ax.hist(protein.missing_pct,bins=np.linspace(0,100,21),color='#527D9E',edgecolor='white',linewidth=.4); ax.set_xlabel('Missing samples per protein group (%)'); ax.set_ylabel('Protein groups'); label(ax,'Protein-level missingness')
ax=missing_axes[2]; good=protein.median_positive_abundance.notna(); xx=np.log10(protein.loc[good,'median_positive_abundance']); yy=protein.loc[good,'missing_pct']
ax.scatter(xx,yy,s=3,alpha=.35,c='#527D9E',edgecolors='none',rasterized=True); ax.set_xlabel('Median observed abundance (log10)'); ax.set_ylabel('Missing samples (%)'); label(ax,'Abundance and missingness')
missing_source = pd.DataFrame(missing[orderp][:,ordered],
                              index=annotation.iloc[orderp,0],
                              columns=sample.iloc[ordered].UniqueSampleID).rename_axis('PG.ProteinGroups').reset_index()
save_series(missing_figures, OUT, ['Figure_02_descriptive_matrix_missingness', 'Figure_02_descriptive_protein_missingness',
                                  'Figure_02_descriptive_abundance_missingness'],
            [missing_source,
             protein[[protein.columns[0],'missing_pct']],
             protein[[protein.columns[0],'median_positive_abundance','missing_pct']]])

def md_table(frame):
    f=frame.copy()
    for c in f.select_dtypes('float'): f[c]=f[c].map(lambda v:f'{v:.2f}')
    return '| '+' | '.join(f.columns)+' |\n| '+' | '.join(['---']*len(f.columns))+' |\n'+'\n'.join('| '+' | '.join(map(str,row))+' |' for row in f.itertuples(index=False,name=None))
cols=['category','samples','detected_any','detected_ge80pct','detected_all','missing_pct','sample_detected_median']
def table_cn(df):
    return md_table(df[cols].rename(columns=dict(zip(cols,['分组','样本数','至少1例检出','至少80%检出','全部样本检出','缺失率(%)','单样本检出中位数']))))
report=f'''# 血浆蛋白组描述性统计

## 数据与口径
原始矩阵共 **{nprot:,} 个蛋白组 × {nsamp} 个样本**。蛋白组ID无空缺、无重复，计数单位为蛋白组，并非拆分后的单个蛋白或基因。
已独立核对原始列头、全部样本位置、metadata一对一对应及原始metadata各字段，均通过。以最终映射表为样本分组依据。实际metadata是工作簿第二张表（名称为Sheet3）。
检出定义为有限且大于0的值。原始空值 {audit['original_null_cells']:,}，零值 {audit['zero_cells']:,}，负值 {audit['negative_cells']:,}，非数值 {audit['non_numeric_cells']:,}，无穷值 {audit['infinite_cells']:,}。未进行填补、过滤、归一化或差异检验。

## 总体
- 至少一个样本检出：**{audit['detected_at_least_once']:,}** 个蛋白组。
- 全部样本均检出：**{audit['detected_in_all_samples']:,}** 个蛋白组。
- 所有样本均未检出：{audit['never_detected']} 个蛋白组。
- 总缺失：**{audit['missing_cells']:,}/{audit['total_cells']:,}（{audit['missing_pct']:.2f}%）**。
- 单样本检出中位数：{sample.detected_protein_groups.median():.0f}；四分位数：{sample.detected_protein_groups.quantile(.25):.0f}–{sample.detected_protein_groups.quantile(.75):.0f}；范围：{sample.detected_protein_groups.min()}–{sample.detected_protein_groups.max()}。

## 不同环境
{table_cn(env)}

两个环境共同检出 {overlap['shared']} 个蛋白组；仅在 {ep.columns[1]} 检出 {overlap[str(ep.columns[1])+'_only']} 个，仅在 {ep.columns[2]} 检出 {overlap[str(ep.columns[2])+'_only']} 个。“仅检出”不等同环境特异表达。

## 不同地区
{table_cn(region)}

地区按原始group字段统计，不推测缩写对应的地理名称。XZ_YD总计28个样本，其中此前区段验证的27个属于特定进样日期。

## 解释边界
组内至少1例检出是检出并集，随样本数变化；80%覆盖阈值用于描述，并非本次过滤标准。缺失率分母为该组全部蛋白组×样本数。
环境与地区在本数据中完全对应：高温包含FJ_FQ、FJ_PT、FJ_QZ、GZ_TH，高应激包含XZ_GG、XZ_YA、XZ_YB、XZ_YC、XZ_YD。因此不能将环境间描述差异直接归因于环境效应。
样本计数是矩阵列数；尚未据受试者ID确认生物学独立重复。unknown/missing暴露标签仍保留在描述中。
丰度图只使用有至少一个有效定量值的蛋白组，完全未检出的{audit['never_detected']}行没有可定义的丰度；缺失模式不能单独证明MAR或MNAR。

## 图注
Figure 01 为四个独立输出：地区样本量（蓝色为高应激、棕色为高温）、环境内检出覆盖、地区内检出覆盖，以及各样本检出深度。样本深度箱体为Q1–Q3，线为中位数，须为1.5×IQR范围内最远值，点为其外观测；全部样本参与。覆盖图为精确计数，无误差条或统计检验。
Figure 02 为三个独立输出：全蛋白组×样本缺失图（蛋白按缺失率排序，样本按环境、地区和原始位置排序；浅色检出、深色缺失）、蛋白缺失率分布，以及检出值中位丰度取log10与缺失率的关系（每点一个蛋白组）。没有插值或缺失填补。

## 可复现性
同目录CSV包括环境、地区、TREAT1和进样日期汇总，蛋白及样本明细、组内蛋白检出率和环境×地区交叉表。audit.json记录输入SHA256及核验计数。图形同时提供PDF、SVG与PNG。
'''
(OUT/'数据描述报告.md').write_text(report,encoding='utf-8')
print(json.dumps(audit,ensure_ascii=False,indent=2),flush=True)
print(env.to_string(index=False),flush=True)
print(region.to_string(index=False),flush=True)
