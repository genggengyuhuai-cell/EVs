"""v2.1: parallel all-protein binary detection branch; no abundance zero filling.

Only main() performs I/O. Frozen upstream scripts are never imported/executed.
The original 70% all-group quantitative filter is untouched.
"""
from pathlib import Path
import hashlib
import importlib.metadata
import json
from datetime import datetime, timezone

import numpy as np
import pandas as pd
from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
ROOT = HERE.parent
VERSION = "2.1"
TARGET_THRESHOLDS = (0.70, 0.60)
OTHER_MAX_EXCLUSIVE = 0.20
GROUPS = ("control", "low", "high")
LABELS = ("Control", "Short", "Long")
RESTRICTED = (
    "Control_specific", "Short_specific", "Long_specific",
    "Control_Short_enriched", "Control_Long_enriched", "Short_Long_enriched",
)


def digest(path, algorithm="sha256"):
    h = hashlib.new(algorithm)
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            h.update(chunk)
    return h.hexdigest()


def unique_ids(series, label):
    if series.isna().any() or series.astype(str).str.strip().eq("").any() or series.duplicated().any():
        raise ValueError(f"{label}: duplicate or empty IDs")


def clean_header(value):
    text = str(value).strip()
    return text[:-2] if text.endswith(".0") else text


def classify_rates(rates, target, other):
    """Specific: one high, both others low; enriched: two high, third low.

    Broad: all three >= target. Sparse: all three < other. Otherwise unbalanced.
    These rules are disjoint when 0 <= other < target <= 1.
    """
    if not 0 <= other < target <= 1:
        raise ValueError("Require 0 <= other < target <= 1")
    high, low = rates >= target, rates < other
    result = np.full(len(rates), "Other_unbalanced", dtype=object)
    result[high.all(axis=1)] = "Broad_detection"
    result[low.all(axis=1)] = "Sparse"
    for i, label in enumerate(LABELS):
        others = [j for j in range(3) if j != i]
        result[high[:, i] & low[:, others].all(axis=1)] = f"{label}_specific"
    for i, j in ((0, 1), (0, 2), (1, 2)):
        k = 3 - i - j
        result[high[:, i] & high[:, j] & low[:, k]] = f"{LABELS[i]}_{LABELS[j]}_enriched"
    return result


def main():
    paths = {
        "source": ROOT / "rawdata" / "processed.xlsx",
        "mapping": ROOT / "rawdata" / "sample_mapping_FINAL.xlsx",
        "audit": HERE / "audit.json",
        "samples": HERE / "sample_statistics.csv",
        "defined": HERE / "dose_defined_metadata.csv",
    }
    for path in paths.values():
        if not path.is_file():
            raise FileNotFoundError(path)
    out = HERE / "detection_pattern"
    if out.exists() and any(out.iterdir()):
        raise FileExistsError(f"Preserve existing outputs; use a new destination: {out}")
    audit = json.loads(paths["audit"].read_text(encoding="utf-8-sig"))
    for key in ("source", "mapping"):
        path = paths[key]
        if digest(path) != audit.get(path.name + "_sha256"):
            raise ValueError(f"Input differs from frozen audit: {path.name}")
    sample = pd.read_csv(paths["samples"], keep_default_na=False)
    required = {"Sheet1_position", "Sheet1_raw_header", "UniqueSampleID", "TREAT1_clean", "condition"}
    if not required.issubset(sample.columns):
        raise ValueError(f"Missing sample columns: {required - set(sample.columns)}")
    sample = sample.sort_values("Sheet1_position").reset_index(drop=True)
    unique_ids(sample.UniqueSampleID, "sample_statistics")
    if sample.Sheet1_position.tolist() != list(range(1, len(sample) + 1)):
        raise ValueError("Sample positions are not a complete 1..N sequence")
    wb = load_workbook(paths["source"], read_only=True, data_only=True)
    try:
        rows = wb.worksheets[0].iter_rows(values_only=True)
        header = next(rows)
        data = pd.DataFrame(list(rows))
    finally:
        wb.close()
    if [clean_header(v) for v in header[7:]] != sample.Sheet1_raw_header.map(clean_header).tolist():
        raise ValueError("Raw matrix headers do not match audited sample-column order")
    annotation = data.iloc[:, :7].copy()
    annotation.columns = header[:7]
    if "PG.ProteinGroups" not in annotation:
        raise ValueError("Missing PG.ProteinGroups in source annotation")
    unique_ids(annotation["PG.ProteinGroups"], "protein groups")
    values = data.iloc[:, 7:].apply(pd.to_numeric, errors="raise").to_numpy(dtype=float)
    if values.shape != (audit["protein_group_rows"], audit["samples"]):
        raise ValueError("Raw matrix dimensions differ from frozen audit")
    keep = sample.TREAT1_clean.isin(GROUPS).to_numpy()
    selected = sample.loc[keep].reset_index(drop=True)
    defined = pd.read_csv(paths["defined"], keep_default_na=False)
    if not {"UniqueSampleID", "TREAT1_clean", "condition"}.issubset(defined):
        raise ValueError("dose_defined_metadata is missing keys")
    unique_ids(defined.UniqueSampleID, "dose_defined_metadata")
    if set(selected.UniqueSampleID) != set(defined.UniqueSampleID):
        raise ValueError("Exposure-defined sample-ID mismatch")
    # Explicit ID match rather than relying on file row order.
    defined = defined.set_index("UniqueSampleID").loc[selected.UniqueSampleID].reset_index()
    for column in ("TREAT1_clean", "condition"):
        if not np.array_equal(selected[column].astype(str), defined[column].astype(str)):
            raise ValueError(f"Metadata conflict: {column}")
    # Preserve any additional available metadata columns for the modelling branch.
    for column in defined.columns:
        if column not in selected:
            selected[column] = defined[column].to_numpy()
    detected = (np.isfinite(values[:, keep]) & (values[:, keep] > 0)).astype(np.uint8)
    table = annotation.copy()  # Full mother table: no protein filtering.
    rates = []
    for group, label in zip(GROUPS, LABELS):
        mask = selected.TREAT1_clean.eq(group).to_numpy()
        n = int(mask.sum())
        if n == 0:
            raise ValueError(f"No samples for exposure group {group}")
        count = detected[:, mask].sum(axis=1)
        table[f"{label}_detected_n"] = count
        table[f"{label}_total_n"] = n
        table[f"{label}_detection_rate"] = count / n
        rates.append(count / n)
    rate_matrix = np.column_stack(rates)
    membership = []
    out.mkdir(parents=True, exist_ok=True)
    table.to_csv(out / "protein_detection_rates_by_exposure.csv", index=False)
    for target in TARGET_THRESHOLDS:
        patterns = classify_rates(rate_matrix, target, OTHER_MAX_EXCLUSIVE)
        tagged = table.assign(Detection_pattern=patterns, Target_threshold=target,
                              Other_threshold_exclusive=OTHER_MAX_EXCLUSIVE)
        tagged.loc[tagged.Detection_pattern.isin(RESTRICTED)].to_csv(
            out / f"exposure_restricted_proteins_{round(100 * target)}pct.csv", index=False)
        membership.append(tagged[["PG.ProteinGroups", "Detection_pattern", "Target_threshold",
                                  "Other_threshold_exclusive"]])
    member = pd.concat(membership, ignore_index=True)
    member.to_csv(out / "detection_pattern_membership.csv", index=False)
    member.groupby(["Target_threshold", "Detection_pattern"]).size().rename("N_proteins").reset_index().to_csv(
        out / "detection_pattern_summary.csv", index=False)
    # Descriptive figures use the same memberships; no extra protein filtering.
    import matplotlib
    matplotlib.use("Agg")
    from nature_plotting import new_figure, save
    for target in TARGET_THRESHOLDS:
        subset = member.loc[member.Target_threshold.eq(target)]
        counts = subset.groupby("Detection_pattern").size().sort_values()
        fig, ax = new_figure(183, max(120, 45 + 7 * len(counts)))
        ax.barh(counts.index.str.replace("_", " "), counts.values, color="#3178A5")
        ax.set(xlabel="Protein groups", title=f"Detection patterns: target {target:.0%}")
        save(fig, out, f"Detection_pattern_counts_{target:.0%}".replace("%", "pct"),
             counts.rename("N_proteins").reset_index())
        for pattern_name in RESTRICTED:
            protein_ids = subset.loc[subset.Detection_pattern.eq(pattern_name), "PG.ProteinGroups"]
            shown = table.loc[table["PG.ProteinGroups"].isin(protein_ids)].copy()
            if shown.empty:
                continue
            columns = [f"{label}_detection_rate" for label in LABELS]
            shown = shown.sort_values(columns, ascending=False)
            fig, ax = new_figure(183, 145)
            im = ax.imshow(shown[columns].to_numpy(), aspect="auto", interpolation="nearest",
                           cmap="Blues", vmin=0, vmax=1)
            ax.set_xticks(range(3), ["Control", "Short exposure", "Long exposure"])
            ax.set_yticks([])
            ax.set(ylabel=f"Protein groups (n = {len(shown)})",
                   title=f"{pattern_name.replace('_', ' ')}: target {target:.0%}")
            fig.colorbar(im, ax=ax, label="Detection proportion", shrink=0.7)
            save(fig, out, f"Detection_rates_{round(target * 100)}pct_{pattern_name}", shown)
    binary = pd.DataFrame(detected, columns=selected.UniqueSampleID)
    binary.insert(0, "PG.ProteinGroups", annotation["PG.ProteinGroups"].to_numpy())
    binary.to_csv(out / "binary_detection_matrix.csv.gz", index=False, compression="gzip")
    selected.to_csv(out / "detection_metadata.csv", index=False)
    sample.loc[~keep, ["UniqueSampleID", "TREAT1_clean"]].to_csv(out / "excluded_exposure_samples.csv", index=False)
    parameters = {
        "script_version": VERSION, "analysis_time_UTC": datetime.now(timezone.utc).isoformat(),
        "detection_definition": "finite quantitative value > 0; binary phenotype only",
        "target_thresholds": list(TARGET_THRESHOLDS), "other_threshold_exclusive": OTHER_MAX_EXCLUSIVE,
        "specific_rule": "one group >= target, both other groups < other",
        "enriched_rule": "two groups >= target, remaining group < other",
        "all_proteins_retained": len(table), "included_samples": len(selected),
        "excluded_samples": int((~keep).sum()),
    }
    for key, path in paths.items():
        parameters[f"{key}_file"] = str(path)
        parameters[f"{key}_MD5"] = digest(path, "md5")
    for package in ("numpy", "pandas", "openpyxl"):
        parameters[f"package_{package}"] = importlib.metadata.version(package)
    pd.DataFrame({"Parameter": list(parameters), "Value": [str(v) for v in parameters.values()]}).to_csv(
        out / "detection_pattern_parameters.csv", index=False)
    (out / "METHODS.txt").write_text(
        "Parallel all-protein detection branch v2.1; not quantitative abundance.\n"
        "No proteins dropped by the quantitative 70% filter. Specific is operational, not absolute uniqueness.\n"
        "Sparse means every group is below 20%; moderate/mixed rates are Other_unbalanced.\n"
        "Neither absence nor zero binary detection establishes biological absence.\n", encoding="utf-8")


if __name__ == "__main__":
    main()
