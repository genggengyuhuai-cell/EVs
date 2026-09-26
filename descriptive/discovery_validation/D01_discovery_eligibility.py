"""Build the Discovery-only protein eligibility universe.

This is a standalone implementation of the frozen Stage 05 eligibility rule.
It intentionally stops after exporting the eligible protein list and the
untransformed Discovery abundance matrix.  It performs no normalization,
imputation, transformation, modelling, contrast testing, FDR calculation, or
DEP analysis.

Nothing is read or written on import.  Execution must be separately authorized.
"""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
import hashlib
import json

import numpy as np
import pandas as pd
from openpyxl import load_workbook


SCRIPT = Path(__file__).resolve()
DV_DIR = SCRIPT.parent
DESCRIPTIVE_DIR = DV_DIR.parent
ROOT = DESCRIPTIVE_DIR.parent

ASSIGNMENT = DESCRIPTIVE_DIR / "discovery_validation_split" / "discovery_validation_assignment.csv"
RAW_WORKBOOK = ROOT / "rawdata" / "processed.xlsx"
MAPPING_WORKBOOK = ROOT / "rawdata" / "sample_mapping_FINAL.xlsx"
SAMPLE_STATISTICS = DESCRIPTIVE_DIR / "sample_statistics.csv"
CANONICAL_ANNOTATION = DESCRIPTIVE_DIR / "canonical_protein_annotation.csv"
SOURCE_AUDIT = DESCRIPTIVE_DIR / "audit.json"
STAGE05_REFERENCE = DESCRIPTIVE_DIR / "05_dose_quantitative_filtering.py"
OUTPUT_DIR = DV_DIR / "D01_discovery_eligibility"

FROZEN_ASSIGNMENT_SHA256 = "062e51026b7420dca2077d5807bfdaa760ac08ac5f19a9af98cf89da2b7b6791"
EXPECTED_ASSIGNMENT_ROWS = 515
EXPECTED_DISCOVERY_N = 386
EXPECTED_VALIDATION_N = 129
EXPECTED_RAW_PROTEINS = 3817
DOSE_LEVELS = ("control", "low", "high")
THRESHOLD_PERCENT = 70
ANNOTATION_COLUMN_COUNT = 7
MAPPING_SHEET = "01_FINAL_MAPPING"

OUTPUT_NAMES = (
    "D01_discovery_eligible_proteins.csv",
    "D01_discovery_eligible_expression.csv.gz",
    "D01_eligibility_diagnostics.csv",
    "D01_participant_mapping_diagnostics.csv",
    "D01_integrity_assertions.csv",
    "D01_manifest.txt",
)

# These are explicitly forbidden as analytical inputs.  The historical Stage 05
# script itself is used only as a provenance reference and is never executed.
PROHIBITED_INPUT_NAMES = {
    "PRIMARY_dose_quantitative_proteins.csv",
    "PRIMARY_dose_quantitative_expression.csv.gz",
    "protein_statistics.csv",
    "dose_protein_detection_rates.csv",
    "dose_filter_membership.csv",
}


def sha256_file(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for block in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(block)
    return digest.hexdigest()


def clean_header(value: object) -> str:
    text = str(value).strip()
    return text[:-2] if text.endswith(".0") else text


class IntegrityChecks:
    """Record every frozen D01 assertion and fail immediately if it is false."""

    def __init__(self) -> None:
        self.records: list[dict[str, str]] = []

    def check(self, assertion_id: str, description: str, condition: bool, detail: str) -> None:
        status = "PASS" if bool(condition) else "FAIL"
        self.records.append(
            {
                "assertion_id": assertion_id,
                "description": description,
                "status": status,
                "detail": detail,
            }
        )
        if status == "FAIL":
            raise AssertionError(f"{assertion_id} FAIL: {description}. {detail}")

    def frame(self) -> pd.DataFrame:
        expected = [f"D01_A{i:02d}" for i in range(1, 29)]
        observed = [record["assertion_id"] for record in self.records]
        if observed != expected:
            raise AssertionError(
                "D01 assertion sequence is incomplete or out of order: "
                f"observed={observed}, expected={expected}"
            )
        return pd.DataFrame(self.records)


def require_files(paths: tuple[Path, ...]) -> None:
    missing = [str(path) for path in paths if not path.is_file()]
    if missing:
        raise FileNotFoundError(f"Required D01 input/reference files are missing: {missing}")


def verify_source_hashes() -> None:
    """Verify locked source files before opening either workbook for data access."""
    audit = json.loads(SOURCE_AUDIT.read_text(encoding="utf-8"))
    expected = {
        RAW_WORKBOOK: audit.get("processed.xlsx_sha256"),
        MAPPING_WORKBOOK: audit.get("sample_mapping_FINAL.xlsx_sha256"),
    }
    for path, expected_hash in expected.items():
        if not expected_hash:
            raise KeyError(f"Source audit has no SHA-256 for {path.name}.")
        observed_hash = sha256_file(path)
        if observed_hash != expected_hash:
            raise AssertionError(
                f"Frozen source hash mismatch for {path.name}: "
                f"observed={observed_hash}; expected={expected_hash}"
            )


def read_and_gate_assignment(checks: IntegrityChecks) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame]:
    """Gate the frozen assignment before any abundance workbook is opened."""
    observed_hash = sha256_file(ASSIGNMENT)
    checks.check(
        "D01_A01",
        "assignment SHA-256 matches the frozen hash",
        observed_hash == FROZEN_ASSIGNMENT_SHA256,
        f"observed_sha256={observed_hash}",
    )

    assignment = pd.read_csv(ASSIGNMENT, keep_default_na=False, dtype={"UniqueSampleID": str})
    required = {"UniqueSampleID", "Split"}
    missing = required.difference(assignment.columns)
    if missing:
        raise ValueError(f"Assignment is missing required columns: {sorted(missing)}")

    duplicate_count = int(assignment["UniqueSampleID"].duplicated().sum())
    checks.check(
        "D01_A02",
        "assignment contains 515 unique participants",
        len(assignment) == EXPECTED_ASSIGNMENT_ROWS
        and assignment["UniqueSampleID"].nunique() == EXPECTED_ASSIGNMENT_ROWS
        and duplicate_count == 0,
        f"rows={len(assignment)}; unique={assignment['UniqueSampleID'].nunique()}; duplicates={duplicate_count}",
    )

    discovery = assignment.loc[assignment["Split"].eq("Discovery")].copy()
    validation = assignment.loc[assignment["Split"].eq("Validation")].copy()
    checks.check(
        "D01_A03",
        "Discovery n equals 386",
        len(discovery) == EXPECTED_DISCOVERY_N,
        f"Discovery_n={len(discovery)}",
    )
    checks.check(
        "D01_A04",
        "Validation n equals 129",
        len(validation) == EXPECTED_VALIDATION_N
        and set(assignment["Split"]) == {"Discovery", "Validation"},
        f"Validation_n={len(validation)}; split_values={sorted(set(assignment['Split']))}",
    )
    return assignment, discovery, validation


def map_discovery_participants(
    discovery: pd.DataFrame,
    validation: pd.DataFrame,
    checks: IntegrityChecks,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Resolve IDs through the frozen metadata/index/position relationship."""
    mapping = pd.read_excel(
        MAPPING_WORKBOOK,
        sheet_name=MAPPING_SHEET,
        dtype={"UniqueSampleID": str},
    )
    sample = pd.read_csv(SAMPLE_STATISTICS, keep_default_na=False, dtype={"UniqueSampleID": str})

    relationship = ["Sheet1_position", "Sheet1_raw_header", "metadata_index", "UniqueSampleID"]
    for name, frame in (("frozen mapping", mapping), ("sample_statistics", sample)):
        missing = set(relationship).difference(frame.columns)
        if missing:
            raise ValueError(f"{name} is missing frozen relationship columns: {sorted(missing)}")

    mapping_rel = mapping[relationship].copy()
    sample_rel = sample[relationship].copy()
    for frame in (mapping_rel, sample_rel):
        frame["Sheet1_position"] = pd.to_numeric(frame["Sheet1_position"], errors="raise").astype(int)
        frame["metadata_index"] = pd.to_numeric(frame["metadata_index"], errors="raise").astype(int)
        frame["Sheet1_raw_header"] = frame["Sheet1_raw_header"].map(clean_header)

    frozen_rel = mapping_rel.merge(
        sample_rel,
        on=relationship,
        how="outer",
        indicator=True,
        validate="one_to_one",
    )
    if not frozen_rel["_merge"].eq("both").all():
        raise AssertionError("Frozen mapping and canonical mapping derivative disagree.")

    selected = discovery[["UniqueSampleID"]].merge(
        sample,
        on="UniqueSampleID",
        how="left",
        validate="one_to_one",
        indicator=True,
    )
    checks.check(
        "D01_A05",
        "exactly 386 abundance columns are selected",
        len(selected) == EXPECTED_DISCOVERY_N and selected["Sheet1_position"].notna().sum() == EXPECTED_DISCOVERY_N,
        f"selected_rows={len(selected)}; mapped_positions={selected['Sheet1_position'].notna().sum()}",
    )

    validation_positions = set(
        mapping_rel.loc[mapping_rel["UniqueSampleID"].isin(validation["UniqueSampleID"]), "Sheet1_position"]
    )
    selected_positions = set(pd.to_numeric(selected["Sheet1_position"], errors="coerce").dropna().astype(int))
    overlap = selected_positions.intersection(validation_positions)
    checks.check(
        "D01_A06",
        "zero Validation abundance columns are selected",
        len(overlap) == 0,
        f"Validation_position_overlap={len(overlap)}",
    )
    checks.check(
        "D01_A07",
        "every Discovery ID maps uniquely to metadata",
        selected["_merge"].eq("both").all()
        and selected["UniqueSampleID"].is_unique
        and sample["UniqueSampleID"].is_unique,
        f"unmatched={int(selected['_merge'].ne('both').sum())}",
    )

    discovery_map = discovery[["UniqueSampleID"]].merge(
        mapping_rel,
        on="UniqueSampleID",
        how="left",
        validate="one_to_one",
        indicator=True,
    )
    checks.check(
        "D01_A08",
        "every Discovery ID maps uniquely to a matrix position",
        discovery_map["_merge"].eq("both").all()
        and discovery_map["Sheet1_position"].notna().all(),
        f"unmapped={int(discovery_map['_merge'].ne('both').sum())}",
    )
    checks.check(
        "D01_A09",
        "selected matrix positions are unique",
        discovery_map["Sheet1_position"].is_unique,
        f"duplicate_positions={int(discovery_map['Sheet1_position'].duplicated().sum())}",
    )

    if "TREAT1_clean" not in selected.columns:
        raise ValueError("sample_statistics.csv is missing TREAT1_clean.")
    checks.check(
        "D01_A10",
        "no Discovery ID is missing from dose metadata",
        selected["TREAT1_clean"].astype(str).str.len().gt(0).all(),
        f"missing_dose={int(selected['TREAT1_clean'].astype(str).str.len().eq(0).sum())}",
    )
    observed_doses = set(selected["TREAT1_clean"])
    dose_counts = selected["TREAT1_clean"].value_counts().reindex(DOSE_LEVELS, fill_value=0)
    checks.check(
        "D01_A11",
        "selected cohort contains only control, low, and high and each group is nonempty",
        observed_doses == set(DOSE_LEVELS) and bool((dose_counts > 0).all()),
        f"doses={sorted(observed_doses)}; counts={dose_counts.to_dict()}",
    )

    selected = selected.drop(columns="_merge")
    selected["Sheet1_position"] = selected["Sheet1_position"].astype(int)
    selected["metadata_index"] = selected["metadata_index"].astype(int)
    return selected, mapping_rel


def read_raw_matrix(selected: pd.DataFrame, checks: IntegrityChecks) -> tuple[pd.DataFrame, np.ndarray, list[str]]:
    """Read all 3,817 raw rows, then select Discovery columns by position."""
    wb = load_workbook(RAW_WORKBOOK, read_only=True, data_only=True)
    try:
        ws = wb.worksheets[0]
        rows = ws.iter_rows(values_only=True)
        header = next(rows)
        data = pd.DataFrame(list(rows))
    finally:
        wb.close()

    if len(header) < ANNOTATION_COLUMN_COUNT:
        raise ValueError("Raw worksheet has fewer than seven annotation columns.")
    annotation = data.iloc[:, :ANNOTATION_COLUMN_COUNT].copy()
    annotation.columns = list(header[:ANNOTATION_COLUMN_COUNT])
    if "PG.ProteinGroups" not in annotation.columns:
        raise ValueError("Raw worksheet does not contain PG.ProteinGroups in its annotation block.")

    protein_ids = annotation["PG.ProteinGroups"]
    checks.check(
        "D01_A12",
        "raw protein universe is 3,817 unique nonmissing PG.ProteinGroups",
        len(annotation) == EXPECTED_RAW_PROTEINS
        and protein_ids.notna().all()
        and protein_ids.astype(str).str.len().gt(0).all()
        and protein_ids.is_unique,
        f"rows={len(annotation)}; missing={int(protein_ids.isna().sum())}; duplicates={int(protein_ids.duplicated().sum())}",
    )

    raw_headers = [clean_header(value) for value in header[ANNOTATION_COLUMN_COUNT:]]
    expected_headers = (
        selected.sort_values("Sheet1_position")["Sheet1_raw_header"].map(clean_header).tolist()
    )
    positions = selected.sort_values("Sheet1_position")["Sheet1_position"].to_numpy(dtype=int)
    zero_based = positions - 1
    if positions.min() < 1 or positions.max() > len(raw_headers):
        raise IndexError("A selected Sheet1_position is outside the raw abundance matrix.")
    actual_headers = [raw_headers[index] for index in zero_based]
    if actual_headers != expected_headers:
        raise AssertionError("Selected raw headers do not match the frozen positional mapping.")

    selected_raw = data.iloc[:, ANNOTATION_COLUMN_COUNT + zero_based].copy()
    values = selected_raw.apply(pd.to_numeric, errors="raise").to_numpy(dtype=float)
    selected_ids = selected.sort_values("Sheet1_position")["UniqueSampleID"].tolist()
    return annotation, values, selected_ids


def apply_frozen_rule(
    annotation: pd.DataFrame,
    raw_values: np.ndarray,
    selected: pd.DataFrame,
    selected_ids: list[str],
    checks: IntegrityChecks,
) -> tuple[pd.DataFrame, pd.DataFrame, pd.DataFrame, np.ndarray]:
    """Apply only finite-and-positive detection and the all-dose 70% rule."""
    detected = np.isfinite(raw_values) & (raw_values > 0)
    sentinel = np.array([[np.nan, np.inf, -np.inf, 0.0, -1.0, 1.0]])
    sentinel_result = np.isfinite(sentinel) & (sentinel > 0)
    checks.check(
        "D01_A13",
        "detection definition is exactly finite AND greater than zero",
        np.array_equal(sentinel_result, np.array([[False, False, False, False, False, True]])),
        "implemented as np.isfinite(values) & (values > 0)",
    )
    checks.check(
        "D01_A14",
        "primary threshold is exactly 70 percent",
        THRESHOLD_PERCENT == 70,
        f"threshold={THRESHOLD_PERCENT}",
    )

    ordered_meta = selected.sort_values("Sheet1_position").reset_index(drop=True)
    per_dose_counts: dict[str, np.ndarray] = {}
    pass_masks: list[np.ndarray] = []
    diagnostic_columns: dict[str, object] = {
        "PG.ProteinGroups": annotation["PG.ProteinGroups"].astype(str).to_numpy()
    }
    evaluated_groups: list[str] = []
    for dose in DOSE_LEVELS:
        group_mask = ordered_meta["TREAT1_clean"].eq(dose).to_numpy()
        n = int(group_mask.sum())
        counts = detected[:, group_mask].sum(axis=1).astype(int)
        passes = counts * 100 >= THRESHOLD_PERCENT * n
        per_dose_counts[dose] = counts
        pass_masks.append(passes)
        evaluated_groups.append(dose)
        diagnostic_columns[f"{dose}_detected_samples"] = counts
        diagnostic_columns[f"{dose}_samples"] = n
        diagnostic_columns[f"{dose}_detection_rate"] = counts / n
        diagnostic_columns[f"{dose}_ge70pct"] = passes

    checks.check(
        "D01_A15",
        "threshold is evaluated separately within each Discovery dose group",
        tuple(evaluated_groups) == DOSE_LEVELS and len(pass_masks) == 3,
        f"evaluated_groups={evaluated_groups}",
    )
    eligible_mask = np.logical_and.reduce(pass_masks)
    checks.check(
        "D01_A16",
        "eligibility requires passing all three dose groups",
        np.array_equal(eligible_mask, pass_masks[0] & pass_masks[1] & pass_masks[2]),
        "combined with np.logical_and.reduce across control/low/high",
    )
    recomputed = np.logical_and.reduce(
        [
            per_dose_counts[dose] * 100
            >= THRESHOLD_PERCENT * int(ordered_meta["TREAT1_clean"].eq(dose).sum())
            for dose in DOSE_LEVELS
        ]
    )
    checks.check(
        "D01_A17",
        "integer comparison is detected_count * 100 >= 70 * n",
        np.array_equal(eligible_mask, recomputed),
        "no floating-point threshold comparison used",
    )
    checks.check(
        "D01_A18",
        "Environment and site do not enter eligibility",
        set(evaluated_groups) == set(DOSE_LEVELS),
        "only TREAT1_clean dose masks index the detection matrix",
    )
    checks.check(
        "D01_A19",
        "no normalization or imputation occurs before eligibility",
        raw_values.shape == detected.shape,
        "eligibility is computed directly from selected raw numeric values",
    )

    analytical_inputs = {
        ASSIGNMENT.name,
        RAW_WORKBOOK.name,
        MAPPING_WORKBOOK.name,
        SAMPLE_STATISTICS.name,
        CANONICAL_ANNOTATION.name,
        SOURCE_AUDIT.name,
    }
    checks.check(
        "D01_A20",
        "historical 1,434 protein list is not an eligibility input",
        "PRIMARY_dose_quantitative_proteins.csv" not in analytical_inputs
        and "PRIMARY_dose_quantitative_expression.csv.gz" not in analytical_inputs,
        f"analytical_inputs={sorted(analytical_inputs)}",
    )
    checks.check(
        "D01_A21",
        "historical full-cohort detection and membership summaries are not inputs",
        PROHIBITED_INPUT_NAMES.isdisjoint(analytical_inputs),
        f"analytical_inputs={sorted(analytical_inputs)}",
    )
    checks.check(
        "D01_A22",
        "no Stage 07, Stage 13A, or canonical 256 DEP input is used",
        not any("stage07" in item.lower() or "stage13" in item.lower() or "256" in item for item in analytical_inputs),
        f"analytical_inputs={sorted(analytical_inputs)}",
    )

    canonical = pd.read_csv(CANONICAL_ANNOTATION, keep_default_na=False)
    if "PG.ProteinGroups" not in canonical.columns or not canonical["PG.ProteinGroups"].is_unique:
        raise ValueError("Canonical annotation must contain unique PG.ProteinGroups.")
    canonical = canonical.set_index("PG.ProteinGroups").reindex(annotation["PG.ProteinGroups"].astype(str))
    if canonical.isna().all(axis=1).any():
        raise ValueError("Canonical annotation does not cover every raw protein group.")
    canonical = canonical.reset_index()

    eligible_proteins = canonical.loc[eligible_mask].copy().reset_index(drop=True)
    for dose in DOSE_LEVELS:
        n = int(ordered_meta["TREAT1_clean"].eq(dose).sum())
        eligible_proteins[f"{dose}_detected_samples"] = per_dose_counts[dose][eligible_mask]
        eligible_proteins[f"{dose}_samples"] = n
        eligible_proteins[f"{dose}_detection_rate"] = per_dose_counts[dose][eligible_mask] / n

    checks.check(
        "D01_A23",
        "output PG.ProteinGroups is nonmissing and unique",
        eligible_proteins["PG.ProteinGroups"].notna().all()
        and eligible_proteins["PG.ProteinGroups"].astype(str).str.len().gt(0).all()
        and eligible_proteins["PG.ProteinGroups"].is_unique,
        f"eligible_rows={len(eligible_proteins)}",
    )
    all_eligible_pass = all(
        np.all(
            per_dose_counts[dose][eligible_mask] * 100
            >= THRESHOLD_PERCENT * int(ordered_meta["TREAT1_clean"].eq(dose).sum())
        )
        for dose in DOSE_LEVELS
    )
    checks.check(
        "D01_A24",
        "every eligible protein independently satisfies all three Discovery thresholds",
        all_eligible_pass,
        "independent post-selection threshold verification",
    )
    checks.check(
        "D01_A25",
        "every output abundance column belongs to Discovery",
        selected_ids == ordered_meta["UniqueSampleID"].tolist()
        and len(selected_ids) == EXPECTED_DISCOVERY_N,
        f"output_sample_columns={len(selected_ids)}",
    )

    eligible_raw = raw_values[eligible_mask, :]
    eligible_quant = np.where(np.isfinite(eligible_raw) & (eligible_raw > 0), eligible_raw, np.nan)
    expression = pd.DataFrame(eligible_quant, columns=selected_ids)
    expression.insert(0, "PG.ProteinGroups", annotation.loc[eligible_mask, "PG.ProteinGroups"].astype(str).to_numpy())
    source_observed = np.isfinite(eligible_raw) & (eligible_raw > 0)
    expression_values = expression.iloc[:, 1:].to_numpy(dtype=float)
    checks.check(
        "D01_A26",
        "output missing positions and observed values match the selected raw matrix exactly",
        np.array_equal(np.isfinite(expression_values), source_observed)
        and np.array_equal(expression_values[source_observed], eligible_raw[source_observed]),
        "non-detected cells remain NA; detected raw values are unchanged",
    )
    checks.check(
        "D01_A27",
        "generated statistics derive solely from Discovery columns",
        raw_values.shape[1] == EXPECTED_DISCOVERY_N
        and detected.shape[1] == EXPECTED_DISCOVERY_N
        and set(selected_ids) == set(selected["UniqueSampleID"]),
        f"statistics_matrix_shape={raw_values.shape}",
    )
    checks.check(
        "D01_A28",
        "execution stops before transformation, normalization, modelling, contrasts, FDR, or DEP",
        True,
        "terminal analytical operation is construction of the untransformed eligible matrix",
    )

    diagnostics = pd.DataFrame(diagnostic_columns)
    diagnostics["eligible_all_three_ge70pct"] = eligible_mask
    return eligible_proteins, expression, diagnostics, eligible_mask


def write_outputs(
    assignment: pd.DataFrame,
    selected: pd.DataFrame,
    eligible_proteins: pd.DataFrame,
    expression: pd.DataFrame,
    diagnostics: pd.DataFrame,
    checks: IntegrityChecks,
) -> None:
    OUTPUT_DIR.mkdir(parents=False, exist_ok=False)

    protein_path = OUTPUT_DIR / OUTPUT_NAMES[0]
    expression_path = OUTPUT_DIR / OUTPUT_NAMES[1]
    diagnostic_path = OUTPUT_DIR / OUTPUT_NAMES[2]
    mapping_path = OUTPUT_DIR / OUTPUT_NAMES[3]
    assertions_path = OUTPUT_DIR / OUTPUT_NAMES[4]
    manifest_path = OUTPUT_DIR / OUTPUT_NAMES[5]

    eligible_proteins.to_csv(protein_path, index=False, encoding="utf-8-sig")
    expression.to_csv(expression_path, index=False, compression="gzip", encoding="utf-8")
    diagnostics.to_csv(diagnostic_path, index=False, encoding="utf-8-sig")

    mapping_diagnostics = selected.sort_values("Sheet1_position").copy()
    mapping_diagnostics.insert(0, "mapping_status", "PASS")
    mapping_diagnostics.to_csv(mapping_path, index=False, encoding="utf-8-sig")
    checks.frame().to_csv(assertions_path, index=False, encoding="utf-8-sig")

    # Verify the serialized abundance matrix still has the exact NA pattern and values.
    saved = pd.read_csv(expression_path)
    original_values = expression.iloc[:, 1:].to_numpy(dtype=float)
    saved_values = saved.iloc[:, 1:].to_numpy(dtype=float)
    if not np.array_equal(np.isnan(saved_values), np.isnan(original_values)):
        raise AssertionError("Serialized expression missing-value positions changed.")
    observed = np.isfinite(original_values)
    if not np.array_equal(saved_values[observed], original_values[observed]):
        raise AssertionError("Serialized expression observed values changed.")

    audit = json.loads(SOURCE_AUDIT.read_text(encoding="utf-8"))
    input_hashes = {
        "assignment": sha256_file(ASSIGNMENT),
        "raw_workbook": sha256_file(RAW_WORKBOOK),
        "mapping_workbook": sha256_file(MAPPING_WORKBOOK),
        "source_audit": sha256_file(SOURCE_AUDIT),
        "stage05_reference": sha256_file(STAGE05_REFERENCE),
        "script": sha256_file(SCRIPT),
    }
    if input_hashes["raw_workbook"] != audit.get("processed.xlsx_sha256"):
        raise AssertionError("Raw workbook hash does not match source audit.")
    if input_hashes["mapping_workbook"] != audit.get("sample_mapping_FINAL.xlsx_sha256"):
        raise AssertionError("Mapping workbook hash does not match source audit.")

    dose_counts = selected["TREAT1_clean"].value_counts().reindex(DOSE_LEVELS, fill_value=0)
    output_hashes = {
        path.name: sha256_file(path)
        for path in (protein_path, expression_path, diagnostic_path, mapping_path, assertions_path)
    }
    lines = [
        "D01 Discovery-only eligibility manifest",
        f"execution_timestamp={datetime.now().astimezone().isoformat()}",
        f"assignment_path={ASSIGNMENT}",
        f"assignment_sha256={input_hashes['assignment']}",
        f"raw_workbook_path={RAW_WORKBOOK}",
        f"raw_workbook_sha256={input_hashes['raw_workbook']}",
        f"mapping_path={MAPPING_WORKBOOK}",
        f"mapping_sha256={input_hashes['mapping_workbook']}",
        f"source_audit_path={SOURCE_AUDIT}",
        f"source_audit_sha256={input_hashes['source_audit']}",
        f"stage05_reference_path={STAGE05_REFERENCE}",
        f"stage05_reference_sha256={input_hashes['stage05_reference']}",
        f"script_path={SCRIPT}",
        f"script_sha256={input_hashes['script']}",
        f"Discovery_n={int(assignment['Split'].eq('Discovery').sum())}",
        f"Validation_n={int(assignment['Split'].eq('Validation').sum())}",
        *(f"Discovery_{dose}_n={int(dose_counts[dose])}" for dose in DOSE_LEVELS),
        f"threshold_percent={THRESHOLD_PERCENT}",
        "detection_rule=np.isfinite(values) & (values > 0)",
        f"raw_protein_count={EXPECTED_RAW_PROTEINS}",
        f"eligible_protein_count={len(eligible_proteins)}",
        "historical_1434_used_as_input=NO",
        "Stage13A_or_canonical256_DEP_used_as_input=NO",
        "Validation_protein_outcomes_accessed=NO",
        *(f"output_sha256[{name}]={digest}" for name, digest in output_hashes.items()),
    ]
    manifest_path.write_text("\n".join(lines) + "\n", encoding="utf-8")


def main() -> None:
    require_files(
        (
            ASSIGNMENT,
            RAW_WORKBOOK,
            MAPPING_WORKBOOK,
            SAMPLE_STATISTICS,
            CANONICAL_ANNOTATION,
            SOURCE_AUDIT,
            STAGE05_REFERENCE,
        )
    )
    if OUTPUT_DIR.exists():
        raise FileExistsError(
            f"Refusing to overwrite existing D01 output directory: {OUTPUT_DIR}"
        )

    checks = IntegrityChecks()
    assignment, discovery, validation = read_and_gate_assignment(checks)
    verify_source_hashes()
    selected, _mapping = map_discovery_participants(discovery, validation, checks)
    annotation, raw_values, selected_ids = read_raw_matrix(selected, checks)
    eligible_proteins, expression, diagnostics, _eligible_mask = apply_frozen_rule(
        annotation, raw_values, selected, selected_ids, checks
    )
    write_outputs(
        assignment,
        selected,
        eligible_proteins,
        expression,
        diagnostics,
        checks,
    )
    print("PASS: D01 Discovery-only eligibility completed.")


if __name__ == "__main__":
    main()
