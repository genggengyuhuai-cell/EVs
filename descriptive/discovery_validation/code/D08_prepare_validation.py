#!/usr/bin/env python3
"""D08 firewall and Validation raw-matrix materialization.

This script must only be invoked as stage D08. It refuses access unless the frozen
D03 lock and an investigator-created unlock marker both exist. It does not select
proteins from Validation and cannot expand the locked family.
"""
from __future__ import annotations
import csv, hashlib, subprocess, sys
from pathlib import Path
import numpy as np
import pandas as pd
from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[3]
BASE = ROOT / "descriptive" / "discovery_validation"
ASSIGN = ROOT / "descriptive" / "discovery_validation_split" / "discovery_validation_assignment.csv"
ASSIGN_SHA = "062e51026b7420dca2077d5807bfdaa760ac08ac5f19a9af98cf89da2b7b6791"

def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()

def clean(x: object) -> str:
    return "" if x is None else str(x).strip()

def main() -> None:
    marker = BASE / "VALIDATION_UNLOCKED.txt"
    lock = BASE / "D03_candidate_lock" / "D03_locked_candidates.csv"
    manifest = BASE / "D03_candidate_lock" / "D03_candidate_lock_manifest.csv"
    if not marker.is_file(): raise SystemExit("VALIDATION_LOCKED: investigator-created VALIDATION_UNLOCKED.txt is required")
    if sha(ASSIGN) != ASSIGN_SHA: raise SystemExit("Frozen assignment hash mismatch")
    m = pd.read_csv(manifest, keep_default_na=False).set_index("Key")["Value"]
    if sha(lock) != m["candidate_list_sha256"]: raise SystemExit("D03 candidate-list hash mismatch")
    candidates = pd.read_csv(lock, usecols=["PG.ProteinGroups"], dtype=str)
    if candidates["PG.ProteinGroups"].duplicated().any(): raise SystemExit("Duplicate locked candidate key")
    a = pd.read_csv(ASSIGN, dtype={"UniqueSampleID": str})
    val = a.loc[a.Split.eq("Validation"), ["UniqueSampleID"]]
    if len(val) != 129 or not val.UniqueSampleID.is_unique: raise SystemExit("Validation assignment invariant failed")
    stats = pd.read_csv(ROOT / "descriptive" / "sample_statistics.csv", dtype={"UniqueSampleID": str})
    sel = val.merge(stats[["UniqueSampleID","Sheet1_position","Sheet1_raw_header"]], on="UniqueSampleID", validate="one_to_one")
    sel = sel.sort_values("Sheet1_position"); positions = sel.Sheet1_position.astype(int).to_numpy()
    wb = load_workbook(ROOT / "rawdata" / "processed.xlsx", read_only=True, data_only=True)
    try:
        rows = wb.worksheets[0].iter_rows(values_only=True); header = next(rows); data = pd.DataFrame(list(rows))
    finally: wb.close()
    proteins = data.iloc[:, 0].astype(str) if header[0] == "PG.ProteinGroups" else data.iloc[:, list(header[:7]).index("PG.ProteinGroups")].astype(str)
    idx = pd.Index(proteins); take = idx.get_indexer(candidates["PG.ProteinGroups"])
    if (take < 0).any(): raise SystemExit("Locked candidate absent from raw protein universe")
    actual = [clean(header[7 + p - 1]) for p in positions]
    if actual != [clean(x) for x in sel.Sheet1_raw_header]: raise SystemExit("Frozen positional mapping mismatch")
    values = data.iloc[take, 7 + positions - 1].apply(pd.to_numeric, errors="raise").to_numpy(float)
    values[~np.isfinite(values) | (values <= 0)] = np.nan
    outdir = BASE / "D08_validation"; outdir.mkdir(parents=True, exist_ok=True)
    outfile = outdir / "D08_validation_locked_raw_expression.csv.gz"
    if outfile.exists(): raise SystemExit(f"Refusing silent overwrite: {outfile}")
    out = pd.DataFrame(values, columns=sel.UniqueSampleID); out.insert(0,"PG.ProteinGroups",candidates["PG.ProteinGroups"]); out.to_csv(outfile,index=False)
    subprocess.run(["Rscript", str(Path(__file__).with_name("D08_validation.R"))], cwd=ROOT, check=True)

if __name__ == "__main__": main()
